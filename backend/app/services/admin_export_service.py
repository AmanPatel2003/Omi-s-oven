import io
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from app.services import admin_analytics_service as analytics


async def _gather_monthly_data(db, date_from: datetime, date_to: datetime) -> dict:
    revenue = await analytics.get_revenue_breakdown(db, date_from, date_to)
    customers = await analytics.get_customer_analytics(db, date_from, date_to)
    rewards = await analytics.get_rewards_analytics(db, date_from, date_to)
    return {"revenue": revenue, "customers": customers, "rewards": rewards}


# ── EXCEL EXPORT ────────────────────────────────────────────────────
async def export_excel(db, date_from: datetime, date_to: datetime) -> io.BytesIO:
    data = await _gather_monthly_data(db, date_from, date_to)

    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="4A4A4A", end_color="4A4A4A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Monthly Report", f"{date_from.date()} to {date_to.date()}"])
    ws.append([])
    ws.append(["Total Revenue", data["revenue"]["total_revenue"]])
    ws.append(["New Customers", data["customers"]["total_new"]])
    ws.append(["Returning Customers", data["customers"]["total_returning"]])
    ws.append(["Points Issued", data["rewards"]["points_issued"]])
    ws.append(["Points Redeemed", data["rewards"]["points_redeemed"]])

    # Sheet 2: Revenue by category
    ws2 = wb.create_sheet("Revenue by Category")
    ws2.append(["Category", "Revenue", "Orders", "% of Total"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    for c in data["revenue"]["by_category"]:
        ws2.append([c["category"], c["revenue"], c["order_count"], c["percentage_of_total"]])

    # Sheet 3: Top products
    ws3 = wb.create_sheet("Top Products")
    ws3.append(["Product", "Revenue", "Units Sold"])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
    for p in data["revenue"]["top_products_by_revenue"]:
        ws3.append([p["name"], p["revenue"], p["units_sold"]])

    # Sheet 4: Customer trend
    ws4 = wb.create_sheet("Customer Trend")
    ws4.append(["Date", "New", "Returning"])
    for cell in ws4[1]:
        cell.fill = header_fill
        cell.font = header_font
    for t in data["customers"]["trend"]:
        ws4.append([t["date"], t["new_customers"], t["returning_customers"]])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ── PDF EXPORT (formatted summary, not a designed report) ─────────────
async def export_pdf(db, date_from: datetime, date_to: datetime) -> io.BytesIO:
    data = await _gather_monthly_data(db, date_from, date_to)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Monthly Business Report", styles["Title"]))
    elements.append(Paragraph(f"{date_from.date()} to {date_to.date()}", styles["Normal"]))
    elements.append(Spacer(1, 20))

    summary_data = [
        ["Metric", "Value"],
        ["Total Revenue", f"₹{data['revenue']['total_revenue']}"],
        ["New Customers", str(data["customers"]["total_new"])],
        ["Returning Customers", str(data["customers"]["total_returning"])],
        ["Points Issued", str(data["rewards"]["points_issued"])],
        ["Points Redeemed", str(data["rewards"]["points_redeemed"])],
    ]
    summary_table = Table(summary_data, colWidths=[250, 200])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A4A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Revenue by Category", styles["Heading2"]))
    cat_data = [["Category", "Revenue", "Orders", "%"]]
    for c in data["revenue"]["by_category"]:
        cat_data.append([c["category"], f"₹{c['revenue']}", str(c["order_count"]), f"{c['percentage_of_total']}%"])
    cat_table = Table(cat_data, colWidths=[150, 100, 100, 100])
    cat_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A4A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements.append(cat_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Top Products", styles["Heading2"]))
    prod_data = [["Product", "Revenue", "Units Sold"]]
    for p in data["revenue"]["top_products_by_revenue"]:
        prod_data.append([p["name"], f"₹{p['revenue']}", str(p["units_sold"])])
    prod_table = Table(prod_data, colWidths=[200, 100, 100])
    prod_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A4A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements.append(prod_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer