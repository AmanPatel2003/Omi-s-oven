import calendar
import io
from datetime import datetime
from bson import ObjectId
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from app.core.exceptions import NotFoundException, BadRequestException, ConflictException

STAFF_ROLES = ["admin", "delivery_staff", "super_admin"]

# ── PAY POLICY — confirm these match your actual rules ──────────────
HALF_DAY_PAY_RATIO = 0.5
LEAVE_IS_PAID = True


def _to_object_id(id_str: str, label: str = "id"):
    try:
        return ObjectId(id_str)
    except Exception:
        raise BadRequestException(f"Invalid {label}")


async def _get_staff_or_404(db, staff_id: str) -> dict:
    oid = _to_object_id(staff_id, "staff id")
    staff = await db.users.find_one({"_id": oid, "role": {"$in": STAFF_ROLES}})
    if not staff:
        raise NotFoundException("Staff member not found")
    return staff


async def _calculate_for_staff(db, staff: dict, month: str) -> dict:
    staff_id = str(staff["_id"])
    details = staff.get("staff_details") or {}
    base_salary = details.get("monthly_salary")

    if base_salary is None:
        raise BadRequestException(f"{staff['name']} has no monthly_salary configured")

    year, mon = map(int, month.split("-"))
    total_days = calendar.monthrange(year, mon)[1]

    cursor = db.attendance.find({"staff_id": staff_id, "date": {"$regex": f"^{month}"}})
    counts = {"present": 0, "half_day": 0, "leave": 0, "absent": 0}
    async for a in cursor:
        if a["status"] in counts:
            counts[a["status"]] += 1

    marked_days = sum(counts.values())
    unmarked_days = total_days - marked_days
    # unmarked days treated as unpaid absence — same as "absent" for pay purposes
    effective_absent = counts["absent"] + unmarked_days

    leave_payable = counts["leave"] if LEAVE_IS_PAID else 0
    payable_days = counts["present"] + (counts["half_day"] * HALF_DAY_PAY_RATIO) + leave_payable

    per_day_rate = base_salary / total_days
    prorated_salary = round(per_day_rate * payable_days, 2)

    existing_record = await db.salary_records.find_one({"staff_id": staff_id, "month": month})
    bonus = existing_record["bonus"] if existing_record else 0
    deductions = existing_record["deductions"] if existing_record else 0
    net_payable = round(prorated_salary + bonus - deductions, 2)

    return {
        "staff_id": staff_id,
        "staff_name": staff["name"],
        "role": staff["role"],
        "month": month,
        "base_salary": base_salary,
        "attendance": {
            "present_days": counts["present"],
            "half_days": counts["half_day"],
            "leave_days": counts["leave"],
            "absent_days": effective_absent,
            "total_working_days": total_days,
        },
        "payable_days": round(payable_days, 2),
        "prorated_salary": prorated_salary,
        "bonus": bonus,
        "deductions": deductions,
        "net_payable": net_payable,
        "already_processed": existing_record is not None,
    }


# ── PREVIEW (dry run, no DB writes) ──────────────────────────────────
async def calculate_month(db, month: str) -> dict:
    cursor = db.users.find({"role": {"$in": STAFF_ROLES}, "is_active": True})
    all_staff = await cursor.to_list(length=None)

    items = []
    for staff in all_staff:
        details = staff.get("staff_details") or {}
        if details.get("monthly_salary") is None:
            continue   # skip staff with no configured salary rather than failing the whole preview
        items.append(await _calculate_for_staff(db, staff, month))

    total_payable = round(sum(i["net_payable"] for i in items), 2)

    return {
        "month": month,
        "total_staff": len(items),
        "total_payable": total_payable,
        "items": items,
    }


# ── PROCESS (finalize + save) ────────────────────────────────────────
async def process_month(db, month: str, admin_id: str, force: bool = False) -> dict:
    preview = await calculate_month(db, month)

    processed_items = []
    skipped_count = 0
    now = datetime.utcnow()

    for calc in preview["items"]:
        if calc["already_processed"] and not force:
            skipped_count += 1
            processed_items.append(calc)
            continue

        doc = {
            "staff_id": calc["staff_id"],
            "month": month,
            "base_salary": calc["prorated_salary"],   # stored base = prorated amount for this month
            "bonus": calc["bonus"],
            "deductions": calc["deductions"],
            "net_paid": calc["net_payable"],
            "payment_date": now,
            "payment_method": "bank_transfer",
            "note": f"Auto-calculated: {calc['attendance']['present_days']} present, "
                    f"{calc['attendance']['half_days']} half-day, {calc['attendance']['leave_days']} leave",
            "recorded_by": admin_id,
            "created_at": now,
        }

        await db.salary_records.update_one(
            {"staff_id": calc["staff_id"], "month": month},
            {"$set": doc},
            upsert=True,
        )
        calc["already_processed"] = True
        processed_items.append(calc)

    total_paid = round(sum(i["net_payable"] for i in processed_items if i["already_processed"]), 2)

    return {
        "month": month,
        "processed_count": len(processed_items) - skipped_count,
        "skipped_count": skipped_count,
        "total_paid": total_paid,
        "items": processed_items,
    }


# ── SINGLE STAFF BREAKDOWN ───────────────────────────────────────────
async def get_staff_salary(db, staff_id: str, month: str) -> dict:
    staff = await _get_staff_or_404(db, staff_id)
    return await _calculate_for_staff(db, staff, month)


# ── GENERATE PDF SLIP ────────────────────────────────────────────────
async def generate_slip_pdf(db, staff_id: str, month: str) -> io.BytesIO:
    calc = await get_staff_salary(db, staff_id, month)
    staff = await _get_staff_or_404(db, staff_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Salary Slip", styles["Title"]))
    elements.append(Paragraph(f"{calc['month']}", styles["Normal"]))
    elements.append(Spacer(1, 20))

    elements.append(Paragraph(f"Employee: {calc['staff_name']}", styles["Normal"]))
    elements.append(Paragraph(f"Role: {calc['role']}", styles["Normal"]))
    elements.append(Paragraph(f"Email: {staff.get('email', '')}", styles["Normal"]))
    elements.append(Spacer(1, 15))

    attendance_data = [
        ["Metric", "Days"],
        ["Present", str(calc["attendance"]["present_days"])],
        ["Half Day", str(calc["attendance"]["half_days"])],
        ["Leave (paid)", str(calc["attendance"]["leave_days"])],
        ["Absent", str(calc["attendance"]["absent_days"])],
        ["Total Days in Month", str(calc["attendance"]["total_working_days"])],
        ["Payable Days", str(calc["payable_days"])],
    ]
    att_table = Table(attendance_data, colWidths=[250, 150])
    att_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A4A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(att_table)
    elements.append(Spacer(1, 20))

    pay_data = [
        ["Component", "Amount (₹)"],
        ["Base Salary (full month)", f"{calc['base_salary']:.2f}"],
        ["Prorated Salary", f"{calc['prorated_salary']:.2f}"],
        ["Bonus", f"{calc['bonus']:.2f}"],
        ["Deductions", f"-{calc['deductions']:.2f}"],
        ["Net Payable", f"{calc['net_payable']:.2f}"],
    ]
    pay_table = Table(pay_data, colWidths=[250, 150])
    pay_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A4A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    elements.append(pay_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


# ── EMAIL SLIP (stub — no provider wired yet) ────────────────────────
async def email_slip(db, staff_id: str, month: str, pdf_buffer: io.BytesIO) -> dict:
    staff = await _get_staff_or_404(db, staff_id)

    from app.core.notification_providers import send_email_with_attachment

    pdf_buffer.seek(0)
    result = await send_email_with_attachment(
        to=staff["email"],
        subject=f"Salary Slip — {month}",
        body="Please find your salary slip attached.",
        pdf_bytes=pdf_buffer.read(),
        filename=f"salary_slip_{month}.pdf",
    )

    return {
        "email_sent": result["success"],
        "email_status_note": None if result["success"] else result.get("error"),
    }


# ── EXPORT ALL SALARIES AS EXCEL ──────────────────────────────────────
async def export_excel(db, month: str) -> io.BytesIO:
    preview = await calculate_month(db, month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Salary {month}"

    header_fill = PatternFill(start_color="4A4A4A", end_color="4A4A4A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = [
        "Staff Name", "Role", "Base Salary", "Present", "Half Day", "Leave", "Absent",
        "Payable Days", "Prorated Salary", "Bonus", "Deductions", "Net Payable", "Status",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for item in preview["items"]:
        ws.append([
            item["staff_name"], item["role"], item["base_salary"],
            item["attendance"]["present_days"], item["attendance"]["half_days"],
            item["attendance"]["leave_days"], item["attendance"]["absent_days"],
            item["payable_days"], item["prorated_salary"], item["bonus"],
            item["deductions"], item["net_payable"],
            "Processed" if item["already_processed"] else "Pending",
        ])

    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "", "", "Total:", preview["total_payable"]])

    for col, width in zip("ABCDEFGHIJKLM", [20, 15, 12, 10, 10, 10, 10, 12, 14, 10, 12, 14, 12]):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer