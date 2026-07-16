import io
import calendar
from datetime import datetime, timedelta
from bson import ObjectId
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.exceptions import NotFoundException, BadRequestException, ConflictException

STAFF_ROLES = ["admin", "delivery_staff", "super_admin"]


def _to_object_id(id_str: str, label: str = "id"):
    try:
        return ObjectId(id_str)
    except Exception:
        raise BadRequestException(f"Invalid {label}")


def _today_str() -> str:
    return datetime.utcnow().strftime("%Y-%m-%d")


async def _get_staff_or_404(db, staff_id: str) -> dict:
    oid = _to_object_id(staff_id, "staff id")
    staff = await db.users.find_one({"_id": oid, "role": {"$in": STAFF_ROLES}})
    if not staff:
        raise NotFoundException("Staff member not found")
    return staff


def _compute_hours(check_in: datetime | None, check_out: datetime | None) -> float | None:
    if check_in and check_out and check_out > check_in:
        return round((check_out - check_in).total_seconds() / 3600, 2)
    return None


def _derive_status(check_in, check_out, hours) -> str:
    if not check_in:
        return "absent"
    if hours is not None and hours < 4:
        return "half_day"
    return "present"


# ── CLOCK IN ──────────────────────────────────────────────────────
async def clock_in(db, staff_id: str, admin_id: str, date: str | None, check_in: datetime | None, note: str | None) -> dict:
    staff = await _get_staff_or_404(db, staff_id)
    date = date or _today_str()
    check_in = check_in or datetime.utcnow()

    existing = await db.attendance.find_one({"staff_id": staff_id, "date": date})
    if existing and existing.get("check_in"):
        raise ConflictException(f"{staff['name']} is already clocked in for {date}")

    now = datetime.utcnow()
    if existing:
        await db.attendance.update_one(
            {"_id": existing["_id"]},
            {"$set": {
                "check_in": check_in, "status": "present", "note": note,
                "marked_by": admin_id, "updated_at": now,
            }},
        )
        record_id = existing["_id"]
    else:
        doc = {
            "staff_id": staff_id, "date": date, "status": "present",
            "check_in": check_in, "check_out": None, "hours_worked": None,
            "note": note, "marked_by": admin_id, "is_corrected": False,
            "created_at": now, "updated_at": now,
        }
        result = await db.attendance.insert_one(doc)
        record_id = result.inserted_id

    return await _out(db, record_id)


# ── CLOCK OUT ─────────────────────────────────────────────────────
async def clock_out(db, staff_id: str, admin_id: str, date: str | None, check_out: datetime | None, note: str | None) -> dict:
    staff = await _get_staff_or_404(db, staff_id)
    date = date or _today_str()
    check_out = check_out or datetime.utcnow()

    existing = await db.attendance.find_one({"staff_id": staff_id, "date": date})
    if not existing or not existing.get("check_in"):
        raise BadRequestException(f"{staff['name']} has not clocked in for {date} yet")
    if existing.get("check_out"):
        raise ConflictException(f"{staff['name']} is already clocked out for {date}")
    if check_out <= existing["check_in"]:
        raise BadRequestException("Clock-out time must be after clock-in time")

    hours = _compute_hours(existing["check_in"], check_out)
    status = _derive_status(existing["check_in"], check_out, hours)

    await db.attendance.update_one(
        {"_id": existing["_id"]},
        {"$set": {
            "check_out": check_out, "hours_worked": hours, "status": status,
            "note": note or existing.get("note"), "marked_by": admin_id,
            "updated_at": datetime.utcnow(),
        }},
    )
    return await _out(db, existing["_id"])


# ── EDIT (CORRECTION) ──────────────────────────────────────────────
async def edit_attendance(db, attendance_id: str, admin_id: str, data: dict) -> dict:
    oid = _to_object_id(attendance_id, "attendance record id")
    record = await db.attendance.find_one({"_id": oid})
    if not record:
        raise NotFoundException("Attendance record not found")

    update_data = {k: v for k, v in data.items() if v is not None}
    if not update_data:
        raise BadRequestException("No fields to update")

    check_in = update_data.get("check_in", record.get("check_in"))
    check_out = update_data.get("check_out", record.get("check_out"))
    if check_in and check_out and check_out <= check_in:
        raise BadRequestException("check_out must be after check_in")

    update_data["hours_worked"] = _compute_hours(check_in, check_out)
    if "status" not in update_data and check_in:
        update_data["status"] = _derive_status(check_in, check_out, update_data["hours_worked"])

    update_data["is_corrected"] = True
    update_data["marked_by"] = admin_id
    update_data["updated_at"] = datetime.utcnow()

    await db.attendance.update_one({"_id": oid}, {"$set": update_data})
    return await _out(db, oid)


# ── MARK LEAVE (date range) ──────────────────────────────────────────
async def mark_leave(db, staff_id: str, admin_id: str, date_from: str, date_to: str, note: str | None) -> list:
    staff = await _get_staff_or_404(db, staff_id)

    start = datetime.strptime(date_from, "%Y-%m-%d")
    end = datetime.strptime(date_to, "%Y-%m-%d")
    if end < start:
        raise BadRequestException("date_to must be on or after date_from")
    if (end - start).days > 60:
        raise BadRequestException("Leave range cannot exceed 60 days in a single request")

    now = datetime.utcnow()
    results = []
    current = start
    while current <= end:
        date_str = current.strftime("%Y-%m-%d")
        await db.attendance.update_one(
            {"staff_id": staff_id, "date": date_str},
            {"$set": {
                "status": "leave", "check_in": None, "check_out": None, "hours_worked": None,
                "note": note, "marked_by": admin_id, "updated_at": now,
            }, "$setOnInsert": {"created_at": now, "is_corrected": False}},
            upsert=True,
        )
        record = await db.attendance.find_one({"staff_id": staff_id, "date": date_str})
        results.append(await _format(db, record))
        current += timedelta(days=1)

    return results


# ── TODAY'S STATUS (all staff) ────────────────────────────────────────
async def get_today_status(db) -> dict:
    date = _today_str()
    staff_cursor = db.users.find({"role": {"$in": STAFF_ROLES}, "is_active": True})
    all_staff = await staff_cursor.to_list(length=None)

    attendance_cursor = db.attendance.find({"date": date})
    attendance_by_staff = {a["staff_id"]: a async for a in attendance_cursor}

    items = []
    present_count = 0
    absent_count = 0
    not_marked_count = 0

    for staff in all_staff:
        sid = str(staff["_id"])
        record = attendance_by_staff.get(sid)
        if record:
            status = record["status"]
            check_in = record.get("check_in")
            check_out = record.get("check_out")
        else:
            status = "not_marked"
            check_in = None
            check_out = None

        if status in ("present", "half_day"):
            present_count += 1
        elif status == "absent":
            absent_count += 1
        elif status == "not_marked":
            not_marked_count += 1

        items.append({
            "staff_id": sid, "staff_name": staff["name"], "role": staff["role"],
            "status": status, "check_in": check_in, "check_out": check_out,
        })

    return {
        "date": date,
        "total_staff": len(all_staff),
        "present_count": present_count,
        "absent_count": absent_count,
        "not_marked_count": not_marked_count,
        "staff": items,
    }


# ── MONTHLY GRID REPORT ──────────────────────────────────────────────
async def get_monthly_report(db, month: str) -> dict:
    year, mon = map(int, month.split("-"))
    days_in_month = calendar.monthrange(year, mon)[1]
    all_dates = [f"{year}-{mon:02d}-{d:02d}" for d in range(1, days_in_month + 1)]

    staff_cursor = db.users.find({"role": {"$in": STAFF_ROLES}, "is_active": True}).sort([("name", 1)])
    all_staff = await staff_cursor.to_list(length=None)

    rows = []
    for staff in all_staff:
        sid = str(staff["_id"])
        cursor = db.attendance.find({"staff_id": sid, "date": {"$regex": f"^{month}"}})
        by_date = {a["date"]: a["status"] async for a in cursor}

        days = [{"date": d, "status": by_date.get(d, "not_marked")} for d in all_dates]

        counts = {"present": 0, "absent": 0, "leave": 0, "half_day": 0}
        for d in days:
            if d["status"] in counts:
                counts[d["status"]] += 1

        rows.append({
            "staff_id": sid, "staff_name": staff["name"], "role": staff["role"],
            "days": days,
            "present_count": counts["present"], "absent_count": counts["absent"],
            "leave_count": counts["leave"], "half_day_count": counts["half_day"],
        })

    return {"month": month, "rows": rows}


# ── EXPORT AS EXCEL ────────────────────────────────────────────────────
async def export_excel(db, month: str) -> io.BytesIO:
    report = await get_monthly_report(db, month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Attendance {month}"

    header_fill = PatternFill(start_color="4A4A4A", end_color="4A4A4A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    status_colors = {
        "present": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "absent": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "leave": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "half_day": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
        "not_marked": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
    }
    status_symbols = {"present": "P", "absent": "A", "leave": "L", "half_day": "H", "not_marked": "-"}

    if not report["rows"]:
        ws.append(["No staff records found for this month"])
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    days_in_month = len(report["rows"][0]["days"])
    header = ["Staff Name", "Role"] + [str(d) for d in range(1, days_in_month + 1)] + ["Present", "Absent", "Leave", "Half Day"]
    ws.append(header)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in report["rows"]:
        row_data = [row["staff_name"], row["role"]]
        row_data += [status_symbols[d["status"]] for d in row["days"]]
        row_data += [row["present_count"], row["absent_count"], row["leave_count"], row["half_day_count"]]
        ws.append(row_data)

        excel_row = ws.max_row
        for col_idx, d in enumerate(row["days"], start=3):
            cell = ws.cell(row=excel_row, column=col_idx)
            cell.fill = status_colors.get(d["status"], status_colors["not_marked"])
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15
    for col in range(3, 3 + days_in_month):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 4

    legend_row = ws.max_row + 2
    ws.cell(row=legend_row, column=1, value="Legend: P=Present, A=Absent, L=Leave, H=Half Day, -=Not Marked")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ── INTERNAL HELPERS ──────────────────────────────────────────────────
async def _format(db, record: dict) -> dict:
    staff = await db.users.find_one({"_id": ObjectId(record["staff_id"])})
    return {
        "id": str(record["_id"]),
        "staff_id": record["staff_id"],
        "staff_name": staff["name"] if staff else "Unknown",
        "date": record["date"],
        "status": record["status"],
        "check_in": record.get("check_in"),
        "check_out": record.get("check_out"),
        "hours_worked": record.get("hours_worked"),
        "note": record.get("note"),
        "is_corrected": record.get("is_corrected", False),
    }


async def _out(db, record_id) -> dict:
    record = await db.attendance.find_one({"_id": record_id})
    return await _format(db, record)